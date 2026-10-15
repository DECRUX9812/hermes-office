from __future__ import annotations

from pathlib import Path

import docx

from conftest import assert_refused_unchanged, call


def _revision_kinds(path: Path) -> list[tuple[str, str]]:
    return [(r.revision_type, r.text) for r in docx.Document(str(path)).revisions]


def test_tracked_edit_writes_real_revisions_and_receipt(files, desk) -> None:
    path = files["docx"]
    out = call("docx_tracked_edit", {"path": str(path), "find": "within 30 days", "replace": "within 45 days"})
    assert out["ok"] is True, out
    assert out["changed"] is True
    assert out["replace"]["tracked"] is True
    kinds = _revision_kinds(path)
    assert ("deletion", "30") in kinds and ("insertion", "45") in kinds
    receipt = desk.get_receipt(out["receipt_id"])
    assert receipt["before_sha256"] != receipt["after_sha256"]
    assert Path(receipt["backup_path"]).read_bytes() != path.read_bytes()
    assert _revision_kinds(Path(receipt["backup_path"])) == []
    ws = desk.get_workstream(out["workstream_id"])
    assert ws["needs_you"] == 1 and ws["status"] == "needs_you"
    assert ws["session_id"] == "session-1"


def test_tracked_edit_refuses_ambiguous_and_missing_targets(files) -> None:
    path = files["docx"]
    out = assert_refused_unchanged("docx_tracked_edit", {"path": str(path), "find": "30 days", "replace": "60 days"}, path, reason="AmbiguousTargetError")
    assert out["refused"] is True and "nth" in out["hint"]
    assert_refused_unchanged("docx_tracked_edit", {"path": str(path), "find": "ninety days", "replace": "x"}, path, reason="TargetNotFoundError")
    ok = call("docx_tracked_edit", {"path": str(path), "find": "30 days", "replace": "60 days", "nth": 1})
    assert ok["ok"] is True, ok


def test_edits_outside_the_desk_are_refused(outside) -> None:
    path = outside["docx"]
    out = assert_refused_unchanged("docx_tracked_edit", {"path": str(path), "find": "within 30 days", "replace": "within 45 days"}, path, reason="outside_desk")
    assert out["refused"] is True


def test_wrong_file_kind_is_refused(files) -> None:
    path = files["xlsx"]
    assert_refused_unchanged("docx_tracked_edit", {"path": str(path), "find": "a", "replace": "b"}, path, reason="unsupported_file")


def test_comment_reply_and_resolve(files) -> None:
    path = files["docx"]
    added = call("docx_comment", {"path": str(path), "find": "1.5% per month", "text": "Is this rate enforceable?", "author": "Reviewer"})
    assert added["ok"] is True, added
    comment_id = added["comment"]["comment_id"]
    replied = call("docx_comment", {"path": str(path), "reply_to": comment_id, "text": "Yes, under clause 2."})
    assert replied["ok"] is True, replied
    resolved = call("docx_comment", {"path": str(path), "resolve": comment_id})
    assert resolved["ok"] is True, resolved
    inspected = call("office_inspect", {"path": str(path), "detail": "full"})
    thread = {c["comment_id"]: c for c in inspected["comments"]}
    assert thread[comment_id]["resolved"] is True
    assert thread[comment_id]["anchored_text"] == "1.5% per month"
    assert any(c["parent_id"] == comment_id for c in thread.values())


def test_comment_refusals_leave_file_unchanged(files) -> None:
    path = files["docx"]
    assert_refused_unchanged("docx_comment", {"path": str(path), "reply_to": 999, "text": "x"}, path, reason="TargetNotFoundError")
    assert_refused_unchanged("docx_comment", {"path": str(path), "find": "not in the document", "text": "x"}, path, reason="TargetNotFoundError")


def test_accept_and_reject_tracked_changes(files, desk) -> None:
    path = files["docx"]
    call("docx_tracked_edit", {"path": str(path), "find": "within 30 days", "replace": "within 45 days"})
    call("docx_tracked_edit", {"path": str(path), "find": "1.5% per month", "replace": "1% per month", "author": "Counsel"})
    only_counsel = call("docx_redline_accept_reject", {"path": str(path), "action": "reject", "scope": "author", "author": "Counsel"})
    assert only_counsel["ok"] is True, only_counsel
    assert only_counsel["resolved"] == 1 and only_counsel["remaining"] == 2
    accepted = call("docx_redline_accept_reject", {"path": str(path), "action": "accept"})
    assert accepted["remaining"] == 0
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    assert "within 45 days" in text and "1.5% per month" in text
    ws = desk.get_workstream(accepted["workstream_id"])
    assert ws["needs_you"] == 0 and ws["status"] == "done"


def test_accept_with_bad_index_is_refused(files) -> None:
    path = files["docx"]
    call("docx_tracked_edit", {"path": str(path), "find": "within 30 days", "replace": "within 45 days"})
    assert_refused_unchanged("docx_redline_accept_reject", {"path": str(path), "action": "accept", "scope": "index", "index": 9}, path, reason="TargetNotFoundError")


def test_compare_builds_a_reviewable_redline(files, desk) -> None:
    original, revised = files["docx"], files["docx_v2"]
    out = call("docx_redline_accept_reject", {"path": str(revised), "action": "compare", "compare_with": str(original), "author": "Hermes"})
    assert out["ok"] is True, out
    redline = Path(out["output_path"])
    assert redline.name == "Services Agreement (Client edits) (redline).docx"
    kinds = _revision_kinds(redline)
    assert any(kind == "insertion" and "45" in text for kind, text in kinds)
    assert any(kind == "deletion" and "30" in text for kind, text in kinds)
    assert _revision_kinds(original) == [] and _revision_kinds(revised) == []
    files_rows = desk.files(out["workstream_id"])
    assert files_rows[0]["state"] == "created"


def test_compare_refuses_inputs_with_pending_revisions(files) -> None:
    original, revised = files["docx"], files["docx_v2"]
    call("docx_tracked_edit", {"path": str(revised), "find": "Contoso Studio LLC", "replace": "Contoso Studio Inc."})
    out = call("docx_redline_accept_reject", {"path": str(revised), "action": "compare", "compare_with": str(original)})
    assert out["ok"] is False and out["reason"] == "UnsupportedStructureError", out
    assert not (revised.parent / "Services Agreement (Client edits) (redline).docx").exists()


def test_search_exact_normalized_and_regex(files) -> None:
    path = files["docx"]
    exact = call("office_search", {"path": str(path), "query": "30 days"})
    assert len(exact["results"][0]["matches"]) == 2
    normalized = call("office_search", {"path": str(path), "query": "PAYMENT IS   due", "match": "normalized"})
    assert len(normalized["results"][0]["matches"]) == 1
    regex = call("office_search", {"path": str(path), "query": r"\d+ days"})
    assert regex["results"][0]["matches"] == []
    regex = call("office_search", {"path": str(path), "query": r"\d+ days", "regex": True})
    assert [m["text"] for m in regex["results"][0]["matches"]] == ["30 days", "30 days"]
    folder = call("office_search", {"folder": str(files["folder"]), "query": "Northwind"})
    assert folder["files_with_matches"] >= 3
