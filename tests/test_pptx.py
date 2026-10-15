from __future__ import annotations

from pptx import Presentation

from conftest import assert_refused_unchanged, call


def _texts(path) -> list[list[str]]:
    return [[s.text_frame.text for s in slide.shapes if s.has_text_frame] for slide in Presentation(str(path)).slides]


def test_inspect_lists_slides_and_shapes(files) -> None:
    out = call("office_inspect", {"path": str(files["pptx"])})
    assert out["ok"] is True and out["summary"]["slides"] == 3
    assert "Revenue" in out["slides"][1]["shapes"]


def test_inspect_formats_reports_provenance(files) -> None:
    out = call("pptx_inspect_formats", {"path": str(files["pptx"]), "slide": 1, "shape": "Revenue"})
    assert out["ok"] is True, out
    font = out["paragraphs"][0]["runs"][0]["font"]
    assert font["size"]["value_pt"] == 20.0
    assert font["name"]["provenance"][-1]["level"].startswith("theme")


def test_replace_text_in_named_shape(files, desk) -> None:
    path = files["pptx"]
    out = call("pptx_edit_shape", {"path": str(path), "slide": 1, "shape": "Revenue", "ops": [{"op": "replace_text", "find": "12%", "replace": "15%"}]})
    assert out["ok"] is True, out
    assert out["applied"][0]["replacements"] == 1
    assert "Revenue grew 15% quarter on quarter" in _texts(path)[1]
    assert desk.get_receipt(out["receipt_id"])["tool"] == "pptx_edit_shape"


def test_shape_and_slide_refusals_leave_deck_unchanged(files) -> None:
    path = files["pptx"]
    assert_refused_unchanged("pptx_edit_shape", {"path": str(path), "slide": 1, "shape": "Nope", "ops": [{"op": "replace_text", "find": "a", "replace": "b"}]}, path, reason="TargetNotFoundError")
    assert_refused_unchanged("pptx_edit_shape", {"path": str(path), "slide": 1, "shape": "Revenue", "ops": [{"op": "replace_text", "find": "99%", "replace": "b"}]}, path, reason="TargetNotFoundError")
    assert_refused_unchanged("pptx_edit_shape", {"path": str(path), "slide": 7, "ops": [{"op": "delete_slide"}]}, path, reason="TargetNotFoundError")


def test_slide_ops_clone_move_delete(files) -> None:
    path = files["pptx"]
    call("pptx_edit_shape", {"path": str(path), "slide": 2, "ops": [{"op": "clone_slide"}]})
    assert len(_texts(path)) == 4
    call("pptx_edit_shape", {"path": str(path), "slide": 3, "ops": [{"op": "move_slide", "to_index": 1}]})
    assert "Next steps" in _texts(path)[1]
    call("pptx_edit_shape", {"path": str(path), "slide": 1, "ops": [{"op": "delete_slide"}]})
    assert len(_texts(path)) == 3


def test_geometry_and_delete_shape(files) -> None:
    path = files["pptx"]
    out = call("pptx_edit_shape", {"path": str(path), "slide": 1, "shape": "Visits", "ops": [{"op": "set_geometry", "x": 914400, "y": 4572000}]})
    assert out["ok"] is True, out
    shape = next(s for s in Presentation(str(path)).slides[1].shapes if s.name == "Visits")
    assert (shape.left, shape.top) == (914400, 4572000)
    call("pptx_edit_shape", {"path": str(path), "slide": 1, "shape": "Visits", "ops": [{"op": "delete_shape"}]})
    assert all(s.name != "Visits" for s in Presentation(str(path)).slides[1].shapes)


def test_compose_imports_slides_into_a_new_deck(files, desk) -> None:
    base, appendix = files["pptx"], files["appendix"]
    output = files["folder"] / "Review with appendix.pptx"
    out = call("pptx_compose_decks", {"path": str(base), "sources": [{"path": str(appendix), "slides": [0]}], "mode": "adopt_theme", "output_path": str(output)})
    assert out["ok"] is True, out
    assert out["slides"] == 4
    assert "Appendix: Method" in _texts(output)[3]
    assert len(_texts(base)) == 3
    assert desk.files(out["workstream_id"])[0]["state"] == "created"


def test_compose_refuses_bad_mode(files) -> None:
    base = files["pptx"]
    out = assert_refused_unchanged("pptx_compose_decks", {"path": str(base), "sources": [{"path": str(files["appendix"])}], "mode": "merge"}, base)
    assert out["reason"] == "bad_request"


def test_diff_reports_text_and_added_shapes(files) -> None:
    out = call("pptx_diff", {"a": str(files["pptx"]), "b": str(files["pptx_v2"])})
    assert out["ok"] is True, out
    changes = out["diff"]["slide_changes"]
    flat = str(changes)
    assert "Revenue grew 15%" in flat and "Hire a regional producer" in flat
