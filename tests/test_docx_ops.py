from __future__ import annotations

import json
from pathlib import Path

from hermes_office.desk import Desk
from hermes_office.docx_ops import accept_or_reject_docx, comment_docx, preview_docx, redline_docx, tracked_edit_docx
from hermes_office.tools import docx_comment, docx_redline_accept_reject, docx_tracked_edit
from hermes_office.guard import sha256
from docx import Document


def test_docx_tracked_edit_success_and_refusal(sample_files: dict[str, Path], desk: Desk, hermes_home: Path) -> None:
    path = sample_files["docx"]
    refusal_path = path.parent / "refusal.docx"
    doc = Document()
    doc.add_paragraph("Payment is due within 30 days of the invoice date.")
    doc.add_paragraph("Payment is due within 30 days of the invoice date.")
    doc.save(refusal_path)
    before = sha256(path)
    ok = json.loads(docx_tracked_edit({"path": str(path), "find": "Payment is due within 30 days of the invoice date.", "replace": "Payment is due within 15 days of the invoice date.", "nth": 1, "author": "A"}))
    assert ok["ok"] is True
    assert sha256(path) != before
    assert desk.list_receipts()

    refusal_before = sha256(refusal_path)
    refused = json.loads(docx_tracked_edit({"path": str(refusal_path), "find": "Payment is due within 30 days of the invoice date.", "replace": "x", "author": "A"}))
    assert refused["refused"] is True
    assert sha256(refusal_path) == refusal_before
    assert any(event["kind"] == "refused" for event in desk.events(desk.list_workstreams()[0]["id"]))


def test_docx_comment_and_redline(sample_files: dict[str, Path], desk: Desk) -> None:
    path = sample_files["docx"]
    result = json.loads(docx_comment({"path": str(path), "find": "Acme Co.", "text": "review", "author": "Editor", "nth": 0}))
    assert result["ok"] is True

    before = sha256(path)
    refused = json.loads(docx_redline_accept_reject({"path": str(path), "action": "accept", "compare_with": str(path.parent / "missing.docx")}))
    assert refused["refused"] is True
    assert sha256(path) == before
    assert any(event["kind"] == "refused" for event in desk.events(desk.list_workstreams()[0]["id"]))


def test_docx_core_helpers(sample_files: dict[str, Path]) -> None:
    preview = preview_docx(sample_files["docx"])
    assert preview["kind"] == "docx"
    assert preview["docx"]["paragraphs"]
    result = redline_docx(sample_files["docx"], action="reject", compare_with=sample_files["docx"])
    assert result["remaining"] >= 0
    assert accept_or_reject_docx(sample_files["docx"], "accept", compare_with=sample_files["docx"])["resolved"] >= 0
    assert comment_docx(sample_files["docx"], find="Acme Co.", text="note", author="A")["comment"]["author"] == "A"
    assert tracked_edit_docx(sample_files["docx"], find="Payment is due within 15 days of the invoice date.", replace="Payment is due within 10 days of the invoice date.")["revisions"] >= 0
