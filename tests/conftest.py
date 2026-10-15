from __future__ import annotations

import os
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from hermes_office.desk import Desk


@pytest.fixture
def hermes_home(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    home = tmp_path / ".hermes"
    home.mkdir()
    monkeypatch.setenv("HERMES_HOME", str(home))
    return home


@pytest.fixture
def data_dir(hermes_home: Path) -> Path:
    data = hermes_home / "plugin-data" / "hermes-office"
    data.mkdir(parents=True, exist_ok=True)
    return data


@pytest.fixture
def desk(data_dir: Path) -> Desk:
    return Desk(data_dir)


def _textbox(slide, left, top, width, height, text, name):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    shape.name = name
    return shape


@pytest.fixture
def sample_files(tmp_path: Path, desk: Desk) -> dict[str, Path]:
    folder = tmp_path / "workspace"
    folder.mkdir()
    desk.add_folder(folder)

    docx_path = folder / "Services Agreement.docx"
    doc = Document()
    doc.add_heading("Services Agreement", level=1)
    doc.add_paragraph("Payment is due within 30 days of the invoice date.")
    doc.add_paragraph("Payment is due within 30 days of the invoice date.")
    doc.add_paragraph("Acme Co. and Northwind agree to the schedule.")
    doc.save(docx_path)

    xlsx_path = folder / "Budget FY26.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = 0.08
    ws["B1"] = 1000
    ws["C1"] = "=B1*(1+A1)"
    ws["A2"] = 5
    ws["B2"] = 7
    ws["C2"] = "=SUM(A2:B2)"
    wb.save(xlsx_path)

    pptx_path = folder / "Quarterly Review.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Acme Co."
    _textbox(slide, Inches(1), Inches(2.5), Inches(4), Inches(1), "Northwind partnership", "hero-box")
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide2, Inches(1), Inches(1), Inches(4), Inches(1), "Revenue grew 12%", "metric-box")
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide3, Inches(1), Inches(1), Inches(4), Inches(1), "Risks and next steps", "risk-box")
    prs.save(pptx_path)

    pptx_v2_path = folder / "Quarterly Review v2.pptx"
    prs2 = Presentation(pptx_path)
    prs2.slides[0].shapes.title.text = "Quarterly Review v2"
    prs2.slides.move(prs2.slides[2], 1)
    _textbox(prs2.slides[1], Inches(1), Inches(3), Inches(4), Inches(1), "New action item", "action-box")
    prs2.save(pptx_v2_path)

    return {
        "folder": folder,
        "docx": docx_path,
        "xlsx": xlsx_path,
        "pptx": pptx_path,
        "pptx_v2": pptx_v2_path,
    }
