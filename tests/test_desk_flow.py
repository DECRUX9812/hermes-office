from __future__ import annotations

from pathlib import Path

import docx

from conftest import assert_refused_unchanged, call
from hermes_office import plugin
from hermes_office.guard import sha256


def test_one_workstream_per_chat_titled_from_the_request(files, desk) -> None:
    plugin._on_turn_start(session_id="chat-a", user_message="Change the payment terms to 45 days and update the budget")
    a = call("docx_tracked_edit", {"path": str(files["docx"]), "find": "within 30 days", "replace": "within 45 days"}, session="chat-a")
    b = call("xlsx_edit", {"path": str(files["xlsx"]), "ops": [{"op": "set", "ref": "Budget!C4", "number": 1}]}, session="chat-a")
    assert a["workstream_id"] == b["workstream_id"]
    ws = desk.get_workstream(a["workstream_id"])
    assert ws["title"] == "Change the payment terms to 45 days and update the budget"
    assert ws["status"] == "needs_you"
    assert {Path(f["path"]).name for f in desk.files(ws["id"])} == {"Services Agreement.docx", "Campaign Budget FY26.xlsx"}
    plugin._on_turn_end(session_id="chat-a")
    assert desk.get_workstream(ws["id"])["status"] == "needs_you"
    other = call("xlsx_edit", {"path": str(files["xlsx"]), "ops": [{"op": "set", "ref": "Budget!C4", "number": 2}]}, session="chat-b")
    assert other["workstream_id"] != ws["id"]


def test_turn_hooks_drive_running_and_done(files, desk) -> None:
    plugin._on_turn_start(session_id="chat-c", user_message="Bump media in Q2")
    out = call("xlsx_edit", {"path": str(files["xlsx"]), "ops": [{"op": "set", "ref": "Budget!C4", "number": 13500}]}, session="chat-c")
    assert desk.get_workstream(out["workstream_id"])["status"] == "running"
    plugin._on_turn_end(session_id="chat-c")
    assert desk.get_workstream(out["workstream_id"])["status"] == "done"


def test_turn_ending_on_a_refusal_hands_over_to_the_user(files, desk) -> None:
    plugin._on_turn_start(session_id="chat-d", user_message="Fix the days")
    ok = call("xlsx_edit", {"path": str(files["xlsx"]), "ops": [{"op": "set", "ref": "Budget!C4", "number": 1}]}, session="chat-d")
    call("docx_tracked_edit", {"path": str(files["docx"]), "find": "30 days", "replace": "60 days"}, session="chat-d")
    plugin._on_turn_end(session_id="chat-d")
    ws = desk.get_workstream(ok["workstream_id"])
    assert ws["status"] == "needs_you" and "refused" in ws["note"]
    assert [e["kind"] for e in desk.events(ws["id"])][-1] == "refused"


def test_explicit_workstream_from_the_desk_page_binds_the_session(files, desk) -> None:
    folder = desk.folder_for_path(files["docx"])
    queued = desk.create_workstream(folder["id"], "Tighten the liability clause", request="...", status="queued")
    out = call("docx_tracked_edit", {"path": str(files["docx"]), "find": "prior 12 months", "replace": "prior 6 months", "workstream_id": queued["id"]}, session="chat-e")
    assert out["workstream_id"] == queued["id"]
    assert desk.get_workstream(queued["id"])["session_id"] == "chat-e"


def test_receipt_revert_restores_exact_bytes(files, desk) -> None:
    path = files["docx"]
    original = sha256(path)
    first = call("docx_tracked_edit", {"path": str(path), "find": "within 30 days", "replace": "within 45 days"})
    call("docx_tracked_edit", {"path": str(path), "find": "1.5% per month", "replace": "1% per month"})
    listed = call("office_receipt", {"action": "list"})
    assert len(listed["receipts"]) == 2
    reverted = call("office_receipt", {"action": "revert", "receipt_id": first["receipt_id"]})
    assert reverted["ok"] is True, reverted
    assert sha256(path) == original
    assert docx.Document(str(path)).revisions.__len__() == 0
    again = call("office_receipt", {"action": "revert", "receipt_id": first["receipt_id"]})
    assert again["ok"] is False and again["reason"] == "not_found"
    shown = call("office_receipt", {"action": "show", "receipt_id": reverted["receipt"]["id"]})
    assert shown["receipt"]["library_report"] == {"reverted": first["receipt_id"]}


def test_reverting_a_created_file_removes_it(files) -> None:
    out = call("pptx_compose_decks", {"path": str(files["pptx"]), "sources": [{"path": str(files["appendix"])}], "mode": "keep_appearance", "output_path": str(files["folder"] / "Combined.pptx")})
    assert (files["folder"] / "Combined.pptx").exists()
    call("office_receipt", {"action": "revert", "receipt_id": out["receipt_id"]})
    assert not (files["folder"] / "Combined.pptx").exists()


def test_desk_actions(files, tmp_path) -> None:
    listed = call("office_desk", {"action": "list"})
    assert [f["name"] for f in listed["folders"]] == ["Client Work"]
    added = call("office_desk", {"action": "add_folder", "path": str(tmp_path)})
    assert added["ok"] is True
    missing = call("office_desk", {"action": "add_folder", "path": str(tmp_path / "nope")})
    assert missing["ok"] is False and missing["reason"] == "not_found"
    call("docx_tracked_edit", {"path": str(files["docx"]), "find": "within 30 days", "replace": "within 45 days"})
    asked = call("office_desk", {"action": "needs_you", "note": "Should the notice period change too?"})
    assert asked["workstream"]["note"] == "Should the notice period change too?"
    done = call("office_desk", {"action": "done", "note": "All set"})
    assert done["workstream"]["status"] == "done" and done["workstream"]["needs_you"] == 0
    status = call("office_desk", {"action": "status"})
    assert status["files"][0]["state"] == "changed"
    removed = call("office_desk", {"action": "remove_folder", "folder_id": added["folder"]["id"]})
    assert removed["ok"] is True


def test_reads_mark_files_working_without_creating_workstreams(files, desk) -> None:
    call("office_inspect", {"path": str(files["pptx"])}, session="chat-f")
    assert desk.list_workstreams() == []
    call("xlsx_edit", {"path": str(files["xlsx"]), "ops": [{"op": "set", "ref": "Budget!C4", "number": 1}]}, session="chat-f")
    call("office_inspect", {"path": str(files["pptx"])}, session="chat-f")
    states = {Path(f["path"]).name: f["state"] for f in desk.files(desk.workstream_for_session("chat-f")["id"])}
    assert states == {"Campaign Budget FY26.xlsx": "changed", "Spring Campaign Review.pptx": "working"}


def test_missing_arguments_are_reported_not_raised(files) -> None:
    path = files["docx"]
    out = assert_refused_unchanged("docx_tracked_edit", {"path": str(path), "find": "x"}, path)
    assert out["reason"] == "bad_request" and "replace" in out["message"]
