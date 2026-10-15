from __future__ import annotations

import pytest
from openpyxl import load_workbook

from conftest import assert_refused_unchanged, call
from hermes_office.xlsx_ops import soffice_path

needs_libreoffice = pytest.mark.skipif(soffice_path() is None, reason="LibreOffice (soffice) is not installed")


def test_inspect_reports_sheets_and_formulas(files) -> None:
    out = call("office_inspect", {"path": str(files["xlsx"])})
    assert out["ok"] is True, out
    names = [s["name"] for s in out["sheets"]]
    assert names == ["Budget", "Assumptions"]
    assert out["sheets"][0]["formula_cells"] == 9


def test_edit_set_formula_and_insert_rows_rewrites_references(files, desk) -> None:
    path = files["xlsx"]
    out = call("xlsx_edit", {"path": str(path), "ops": [
        {"op": "set", "ref": "Budget!C4", "number": 13500},
        {"op": "insert_rows", "sheet": "Budget", "index": 6, "amount": 1},
        {"op": "set", "ref": "Budget!A6", "value": "Events"},
        {"op": "set", "ref": "Budget!B6", "number": 1000},
        {"op": "formula", "ref": "Budget!F6", "value": "SUM(B6:E6)"},
    ]})
    assert out["ok"] is True, out
    assert "xl/worksheets/sheet1.xml" in out["edit_receipt"]["cells_changed"]
    ws = load_workbook(str(path))["Budget"]
    assert ws["C4"].value == 13500
    assert ws["A7"].value == "Total"
    assert ws["B7"].value == "=SUM(B2:B5)"
    assert load_workbook(str(path))["Assumptions"]["B3"].value == "=Budget!F7*(1+B2)"
    assert desk.get_receipt(out["receipt_id"])["library_report"]["edit_receipt"]["schema"] == "edit_receipt"


def test_edit_refusals_leave_workbook_unchanged(files) -> None:
    path = files["xlsx"]
    assert_refused_unchanged("xlsx_edit", {"path": str(path), "ops": [{"op": "set", "ref": "Nope!A1", "value": "x"}]}, path, reason="TargetNotFoundError")
    assert_refused_unchanged("xlsx_edit", {"path": str(path), "ops": [{"op": "set", "ref": "Budget!A1", "value": "ok"}, {"op": "explode"}]}, path, reason="bad_request")


@needs_libreoffice
def test_recalc_refreshes_cached_values_with_libreoffice(files) -> None:
    path = files["xlsx"]
    call("xlsx_edit", {"path": str(path), "ops": [{"op": "set", "ref": "Budget!B2", "number": 20000}]})
    out = call("xlsx_recalc", {"path": str(path), "read": ["Budget!F2", "Budget!F6", "Assumptions!B3"]})
    assert out["ok"] is True, out
    assert out["recalculated"] is True
    assert out["values"]["Budget!F2"] == 43000
    assert out["values"]["Budget!F6"] == 164000
    assert out["values"]["Assumptions!B3"] == pytest.approx(164000 * 1.08)
    assert load_workbook(str(path))["Budget"]["F2"].value == "=SUM(B2:E2)"
    assert not [p for p in path.parent.iterdir() if p.name.startswith(".hermes-office-recalc-")]


def test_diff_reports_changed_cells(files, tmp_path) -> None:
    path = files["xlsx"]
    before = tmp_path / "before.xlsx"
    before.write_bytes(path.read_bytes())
    call("xlsx_edit", {"path": str(path), "ops": [{"op": "set", "ref": "Budget!C4", "number": 13500}]})
    out = call("xlsx_diff", {"a": str(before), "b": str(path)})
    assert out["ok"] is True, out
    assert {"address": "Budget!C4", "before": 12000, "after": 13500} == {k: out["diff"]["changed"][0][k] for k in ("address", "before", "after")}


def test_legacy_formats_are_refused(files) -> None:
    legacy = files["folder"] / "old.xls"
    legacy.write_bytes(b"not really a workbook")
    assert_refused_unchanged("xlsx_edit", {"path": str(legacy), "ops": [{"op": "set", "ref": "A1", "value": "x"}]}, legacy, reason="unsupported_file")
