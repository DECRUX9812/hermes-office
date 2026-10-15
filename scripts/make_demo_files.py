"""Generate the demo/test documents with the Paper libraries themselves (no binaries are checked in)."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches, Pt

AGREEMENT = [
    ("Heading 1", "Services Agreement"),
    (None, "This Services Agreement is made between Northwind Traders Ltd. (the Client) and Contoso Studio LLC (the Studio)."),
    ("Heading 2", "1. Services"),
    (None, "The Studio will design and deliver the spring campaign described in Schedule A."),
    ("Heading 2", "2. Fees and payment"),
    (None, "The Client will pay the fees in Schedule B. Payment is due within 30 days of the invoice date."),
    (None, "Late payments accrue interest at 1.5% per month."),
    ("Heading 2", "3. Liability"),
    (None, "Each party's total liability is capped at the fees paid in the prior 12 months."),
    ("Heading 2", "4. Term"),
    (None, "This agreement starts on 1 March 2026 and ends on 28 February 2027."),
    (None, "Either party may end this agreement with 30 days written notice."),
]


def make_docx(path: Path, *, revised: bool = False) -> Path:
    doc = Document()
    for style, text in AGREEMENT:
        if revised:
            text = text.replace("within 30 days of the invoice date", "within 45 days of the invoice date").replace("1.5% per month", "1% per month")
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Milestone"
    table.rows[0].cells[1].text = "Fee (USD)"
    for name, fee in (("Concepts", "12,000"), ("Production", "28,000"), ("Launch support", "6,000")):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, fee
    doc.save(str(path))
    return path


def make_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Line item", "Q1", "Q2", "Q3", "Q4", "Total"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    rows = [("Design", 12000, 9000, 7000, 7000), ("Production", 15000, 18000, 16000, 14000), ("Media", 8000, 12000, 14000, 16000), ("Contingency", 2000, 2000, 2000, 2000)]
    for index, (name, *quarters) in enumerate(rows, start=2):
        ws.append([name, *quarters, f"=SUM(B{index}:E{index})"])
    total_row = len(rows) + 2
    ws.append(["Total", *[f"=SUM({col}2:{col}{total_row - 1})" for col in "BCDEF"]])
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 16
    assumptions = wb.create_sheet("Assumptions")
    assumptions.append(["Assumption", "Value"])
    assumptions.append(["Growth", 0.08])
    assumptions.append(["Next year total", f"=Budget!F{total_row}*(1+B2)"])
    wb.save(str(path))
    return path


def _textbox(slide, x: float, y: float, w: float, h: float, text: str, name: str, size: int = 20) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    box.text_frame.text = text
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def make_pptx(path: Path) -> Path:
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Spring Campaign Review"
    title.placeholders[1].text = "Northwind Traders · Q1 2026"
    results = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(results, 0.6, 0.5, 8.8, 0.9, "Results", "Heading", 32)
    _textbox(results, 0.6, 1.6, 4.2, 1.2, "Revenue grew 12% quarter on quarter", "Revenue")
    _textbox(results, 5.2, 1.6, 4.2, 1.2, "Store visits up 8%", "Visits")
    nxt = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(nxt, 0.6, 0.5, 8.8, 0.9, "Next steps", "Heading", 32)
    _textbox(nxt, 0.6, 1.6, 8.8, 2.0, "Extend the campaign to two new regions", "Plan")
    prs.save(str(path))
    return path


def make_pptx_v2(path: Path, source: Path) -> Path:
    prs = Presentation(str(source))
    slides = list(prs.slides)
    for shape in slides[1].shapes:
        if shape.name == "Revenue":
            shape.text_frame.paragraphs[0].runs[0].text = "Revenue grew 15% quarter on quarter"
    _textbox(slides[2], 0.6, 3.8, 8.8, 1.0, "Hire a regional producer by May", "Owner")
    prs.save(str(path))
    return path


def make_appendix(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Appendix: Method"
    _textbox(slide, 0.6, 1.8, 8.8, 1.5, "Survey of 1,200 shoppers across 14 stores", "Method")
    prs.save(str(path))
    return path


def make_all(folder: Path) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    deck = make_pptx(folder / "Spring Campaign Review.pptx")
    return {
        "docx": make_docx(folder / "Services Agreement.docx"),
        "docx_v2": make_docx(folder / "Services Agreement (Client edits).docx", revised=True),
        "xlsx": make_xlsx(folder / "Campaign Budget FY26.xlsx"),
        "pptx": deck,
        "pptx_v2": make_pptx_v2(folder / "Spring Campaign Review v2.pptx", deck),
        "appendix": make_appendix(folder / "Appendix.pptx"),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output_dir")
    out = Path(parser.parse_args().output_dir).expanduser().resolve()
    for name, path in make_all(out).items():
        print(f"{name}: {path}")


if __name__ == "__main__":
    main()
