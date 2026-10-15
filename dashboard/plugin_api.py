from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException


def _load_core():
    try:
        import hermes_office as core  # type: ignore

        return core
    except Exception:
        root = Path(__file__).resolve().parents[1]
        spec = importlib.util.spec_from_file_location(
            "hermes_office",
            root / "hermes_office" / "__init__.py",
            submodule_search_locations=[str(root / "hermes_office")],
        )
        if spec is None or spec.loader is None:
            raise
        module = importlib.util.module_from_spec(spec)
        sys.modules["hermes_office"] = module
        spec.loader.exec_module(module)
        return module


core = _load_core()
from hermes_office.bridge import conversation, data_dir
from hermes_office.desk import Desk
from hermes_office.docx_ops import redline_docx
from hermes_office.diff_ops import diff_file
from hermes_office.guard import atomic_write, backup, kind_of, resolve_path, sha256
from hermes_office.preview import preview_file

router = APIRouter()


def _desk() -> Desk:
    return Desk(data_dir())


def _file_rows(folder_path: Path) -> list[dict[str, Any]]:
    rows = []
    desk = _desk()
    workstreams = desk.list_workstreams()
    for child in sorted(folder_path.rglob("*")):
        if len(child.relative_to(folder_path).parts) > 4:
            continue
        if any(part.startswith(".") for part in child.relative_to(folder_path).parts[:-1]):
            continue
        if child.is_file() and child.name.startswith("~$"):
            continue
        if child.is_file() and child.suffix.lower() in {".docx", ".pptx", ".xlsx"}:
            count = sum(1 for ws in workstreams if any(file_row["path"] == str(child.resolve()) for file_row in desk.files(ws["id"])))
            rows.append({"path": str(child), "rel": str(child.relative_to(folder_path)), "kind": kind_of(child), "size": child.stat().st_size, "mtime": child.stat().st_mtime, "workstreams": count})
        if len(rows) >= 500:
            break
    return rows


@router.get("/health")
def health() -> dict[str, Any]:
    from openpyxl import oracle

    return {
        "ok": True,
        "version": "0.1.0",
        "engines": {
            "docx": True,
            "pptx": True,
            "xlsx": True,
        },
        "libreoffice": oracle.find_soffice() is not None,
        "data_dir": str(data_dir()),
    }


@router.get("/desk")
def get_desk() -> dict[str, Any]:
    desk = _desk()
    folders = desk.list_folders()
    return {"folders": folders, "needs_you": [ws for ws in desk.list_workstreams() if ws.get("needs_you")], "recent": desk.list_workstreams()}


@router.post("/folders")
def add_folder(payload: dict[str, Any]) -> dict[str, Any]:
    path = resolve_path(payload["path"])
    if not path.exists() or not path.is_dir():
        raise HTTPException(status_code=400, detail="folder does not exist")
    return _desk().add_folder(path)


@router.delete("/folders/{folder_id}")
def delete_folder(folder_id: int) -> dict[str, Any]:
    _desk().remove_folder(folder_id)
    return {"ok": True}


@router.get("/folders/{folder_id}/files")
def folder_files(folder_id: int) -> list[dict[str, Any]]:
    desk = _desk()
    folder = next((item for item in desk.list_folders() if item["id"] == folder_id), None)
    if folder is None:
        raise HTTPException(status_code=404, detail="folder not found")
    return _file_rows(Path(folder["path"]))


@router.get("/folders/{folder_id}/workstreams")
def folder_workstreams(folder_id: int) -> list[dict[str, Any]]:
    return _desk().list_workstreams(folder_id=folder_id)


@router.post("/workstreams")
def create_workstream(payload: dict[str, Any]) -> dict[str, Any]:
    desk = _desk()
    folder_id = int(payload["folder_id"])
    title = payload["title"]
    return desk.create_workstream(folder_id, title, request=payload.get("request"), session_id=payload.get("session_id"), status="queued")


@router.patch("/workstreams/{workstream_id}")
def patch_workstream(workstream_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    return _desk().update_workstream(workstream_id, **payload)


@router.get("/workstreams/{workstream_id}")
def get_workstream(workstream_id: int) -> dict[str, Any]:
    desk = _desk()
    workstream = desk.get_workstream(workstream_id)
    if workstream is None:
        raise HTTPException(status_code=404, detail="workstream not found")
    return {"workstream": workstream, "now": desk.now(workstream_id), "files": desk.files(workstream_id), "events": desk.events(workstream_id), "receipts": desk.list_receipts(workstream_id)}


@router.get("/workstreams/{workstream_id}/conversation")
def workstream_conversation(workstream_id: int) -> dict[str, Any]:
    desk = _desk()
    ws = desk.get_workstream(workstream_id)
    if ws is None:
        raise HTTPException(status_code=404, detail="workstream not found")
    if not ws.get("session_id"):
        return {"available": False, "reason": "missing_session_id"}
    return conversation(ws["session_id"])


@router.get("/preview")
def preview(path: str) -> dict[str, Any]:
    resolved = resolve_path(path)
    if not resolved.exists():
        raise HTTPException(status_code=404, detail="file not found")
    return preview_file(resolved)


@router.get("/diff")
def diff(path: str, receipt_id: int | None = None) -> dict[str, Any]:
    resolved = resolve_path(path)
    if not resolved.exists():
        raise HTTPException(status_code=404, detail="file not found")
    desk = _desk()
    receipt = desk.get_receipt(receipt_id) if receipt_id is not None else next((item for item in desk.list_receipts() if item["path"] == str(resolved)), None)
    if receipt is None:
        raise HTTPException(status_code=404, detail="receipt not found")
    baseline = receipt.get("backup_path")
    if not baseline:
        raise HTTPException(status_code=400, detail="receipt has no backup")
    return diff_file(resolved, baseline)


@router.post("/redline")
def redline(payload: dict[str, Any]) -> dict[str, Any]:
    try:
        desk = _desk()
        path = resolve_path(payload["path"])
        if not path.exists():
            raise HTTPException(status_code=404, detail="file not found")
        folder = desk.folder_for_path(path)
        if folder is None:
            raise HTTPException(status_code=400, detail="outside desk")
        ws = desk.create_workstream(folder["id"], "docx_redline_accept_reject", request="redline", session_id=payload.get("session_id"), status="running")
        before = sha256(path)
        backup_path = backup(desk.data_dir, ws["id"], "redline", path)
        result = redline_docx(path, action=payload["action"], author=payload.get("author"), compare_with=payload.get("compare_with"))
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "docx_redline_accept_reject", path, payload["action"], before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed", receipt_id=receipt["id"], kind="docx")
        return {"ok": True, "resolved": result["resolved"], "receipt": receipt}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/receipts/{receipt_id}/revert")
def revert_receipt(receipt_id: int) -> dict[str, Any]:
    desk = _desk()
    receipt = desk.get_receipt(receipt_id)
    if receipt is None:
        raise HTTPException(status_code=404, detail="receipt not found")
    backup_path = receipt.get("backup_path")
    if backup_path is None:
        raise HTTPException(status_code=400, detail="receipt has no backup")
    target = Path(receipt["path"])
    revert_backup = backup(desk.data_dir, receipt["workstream_id"], f"revert-{receipt_id}", target)
    atomic_write(target, lambda tmp: tmp.write_bytes(Path(backup_path).read_bytes()))
    new_receipt = desk.add_receipt(receipt["workstream_id"], "office_receipt", target, f"revert receipt {receipt_id}", None, None, revert_backup, {"ok": True})
    desk.mark_reverted(receipt_id)
    return {"ok": True, "receipt": new_receipt}


@router.get("/events")
def events(since: str | None = None) -> list[dict[str, Any]]:
    return _desk().events_since(since)
