from __future__ import annotations

from pathlib import Path
from typing import Any

from docx.search import find_text
from openpyxl import load_workbook
from pptx.inspect import inspect_text

from .guard import kind_of, resolve_path
from .docx_ops import _load as load_docx


def _norm(text: str) -> str:
    return " ".join(text.lower().split())


def _search_one(path: str | Path, query: str, *, match: str = "exact", regex: bool = False) -> list[dict[str, Any]]:
    kind = kind_of(path)
    if kind == "docx":
        doc = load_docx(path)
        if regex:
            spans = find_text(doc, query)
        else:
            spans = find_text(doc, query, match=match)
        return [span.to_dict() if hasattr(span, "to_dict") else {"text": getattr(span, "text", query)} for span in spans]
    if kind == "pptx":
        from .pptx_ops import _load as load_pptx

        prs = load_pptx(path)
        hits = []
        for slide_index, slide in enumerate(prs.slides):
            inspection = inspect_text(slide).to_dict()
            for block in inspection.get("blocks", []):
                text = block.get("text", "")
                hay = text if match == "exact" else _norm(text)
                needle = query if match == "exact" else _norm(query)
                if (needle in hay) if regex else (needle == hay or needle in hay):
                    hits.append({"slide": slide_index, "shape": block.get("shape_name"), "text": text})
        return hits
    if kind == "xlsx":
        wb = load_workbook(str(resolve_path(path)))
        return wb.search(query, regex=regex)
    raise ValueError(f"unsupported kind: {kind}")


def search_path_or_folder(path: str | Path | None = None, folder: str | Path | None = None, *, query: str, match: str = "exact", regex: bool = False) -> list[dict[str, Any]]:
    target = folder or path
    if target is None:
        raise ValueError("path or folder is required")
    resolved = resolve_path(target)
    if resolved.is_dir():
        hits = []
        for child in sorted(resolved.rglob("*")):
            if child.is_file() and child.suffix.lower() in {".docx", ".pptx", ".xlsx"}:
                hits.append({"path": str(child), "matches": _search_one(child, query, match=match, regex=regex)})
        return hits
    return [{"path": str(resolved), "matches": _search_one(resolved, query, match=match, regex=regex)}]
