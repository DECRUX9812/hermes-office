from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook, oracle
from openpyxl.errors import OracleTimeoutError
from openpyxl.preserve import diff_workbooks, scan_errors

from .guard import resolve_path

_AUTO = object()


def _load(path: str | Path, *, data_only: bool = False):
    return load_workbook(str(resolve_path(path)), data_only=data_only)


def inspect_xlsx(path: str | Path, *, detail: str = "summary") -> dict[str, Any]:
    wb = _load(path)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({"name": ws.title, "dims": ws.calculate_dimension(), "formulas": sum(1 for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))})
    return {"kind": "xlsx", "path": str(resolve_path(path)), "workbook": {"sheets": sheets, "defined_names": list(getattr(wb, "defined_names", {}).keys()) if hasattr(getattr(wb, "defined_names", {}), "keys") else []}}


def edit_xlsx(path: str | Path, *, ops: list[dict[str, Any]]) -> dict[str, Any]:
    wb = _load(path)
    for op in ops:
        kind = op.get("kind") or op.get("op")
        sheet = wb[op.get("sheet") or wb.sheetnames[0]]
        if kind == "set":
            sheet[op["ref"]] = op["value"]
        elif kind == "formula":
            formula = op["value"]
            sheet[op["ref"]] = formula if str(formula).startswith("=") else f"={formula}"
        elif kind == "insert_rows":
            sheet.insert_rows(int(op["index"]), int(op.get("amount", 1)))
        elif kind == "delete_rows":
            sheet.delete_rows(int(op["index"]), int(op.get("amount", 1)))
        elif kind == "insert_cols":
            sheet.insert_cols(int(op["index"]), int(op.get("amount", 1)))
        elif kind == "delete_cols":
            sheet.delete_cols(int(op["index"]), int(op.get("amount", 1)))
        elif kind == "rename_sheet":
            sheet.title = op["new_name"]
        else:
            raise ValueError(f"unknown xlsx op: {kind}")
    target = resolve_path(path)
    tmp = target.parent / f".{target.name}.edited.xlsx"
    receipt = wb.save(str(tmp), receipt=True)
    tmp.replace(target)
    return {"receipt": receipt.to_dict()}


def recalc_xlsx(path: str | Path, *, read: list[str], soffice: str | None | object = _AUTO) -> dict[str, Any]:
    wb = _load(path)
    if soffice is None:
        return {"ok": True, "recalculated": False, "reason": "libreoffice_missing", "errors": scan_errors(wb)}
    soffice_path = soffice if soffice is not _AUTO else oracle.find_soffice()
    if soffice_path is None:
        return {"ok": True, "recalculated": False, "reason": "libreoffice_missing", "errors": scan_errors(wb)}
    target = resolve_path(path)
    tmp = target.parent / f".{target.name}.recalc.xlsx"
    result = oracle.recalc(str(target), output_path=str(tmp))
    tmp.replace(target)
    values = {}
    data_wb = _load(path, data_only=True)
    for ref in read:
        sheet_name, cell_ref = ref.split("!")
        values[ref] = data_wb[sheet_name][cell_ref].value
    return {"ok": True, "recalculated": True, "values": values, "result": result.to_dict()}


def diff_xlsx(path_a: str | Path, path_b: str | Path) -> dict[str, Any]:
    return diff_workbooks(str(resolve_path(path_a)), str(resolve_path(path_b))).to_dict()


def preview_xlsx(path: str | Path) -> dict[str, Any]:
    wb = _load(path)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(60, ws.max_row or 1), min_col=1, max_col=min(20, ws.max_column or 1)):
            row_data = []
            for cell in row:
                value = cell.value
                row_data.append(
                    {
                        "ref": cell.coordinate,
                        "value": value,
                        "formula": value if isinstance(value, str) and value.startswith("=") else None,
                    }
                )
            rows.append(row_data)
        sheets.append({"name": ws.title, "dims": ws.calculate_dimension(), "rows": rows})
    return {"kind": "xlsx", "path": str(resolve_path(path)), "xlsx": {"sheets": sheets}}
