from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import docx
from docx.commentops import comment_thread, reply, resolve
from docx.errors import TargetNotFoundError
from docx.package import compare, patch_save
from docx.search import find_one, find_text

from .guard import resolve_path

DEFAULT_AUTHOR = "Hermes"

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _load(path: str | Path) -> docx.Document:
    return docx.Document(str(resolve_path(path)))


def _save_in_place(path: str | Path, document: docx.Document) -> dict[str, Any]:
    # patch_save keeps unrelated parts byte-identical and writes atomically (temp file + rename).
    target = str(resolve_path(path))
    result = patch_save(target, document, target)
    return result.to_dict() if hasattr(result, "to_dict") else {}


def _span_dict(span) -> dict[str, Any]:
    return {
        "text": span.text,
        "story": span.story,
        "anchor": span.anchor.to_dict(),
        "in_insert": span.in_insert,
        "in_delete": span.in_delete,
        "crosses_paragraphs": span.crosses_paragraphs,
    }


def inspect_docx(path: str | Path, *, detail: str = "summary") -> dict[str, Any]:
    document = _load(path)
    revisions = document.revisions
    paragraphs = [
        {"index": index, "style": paragraph.style.name if paragraph.style is not None else None, "text": paragraph.text}
        for index, paragraph in enumerate(document.paragraphs)
    ]
    result: dict[str, Any] = {
        "kind": "docx",
        "path": str(resolve_path(path)),
        "summary": {
            "paragraphs": len(paragraphs),
            "tables": len(document.tables),
            "sections": len(document.sections),
            "comments": len(comment_thread(document)),
            "revisions": len(revisions),
            "unsupported_revisions": revisions.remaining_unsupported(),
        },
        "headings": [p for p in paragraphs if (p["style"] or "").lower().startswith(("heading", "title"))],
    }
    if detail == "full":
        result["paragraphs"] = paragraphs
        result["tables"] = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
        result["comments"] = list(comment_thread(document))
        result["revisions"] = [revision.to_dict() for revision in revisions]
    else:
        result["paragraphs_preview"] = [p for p in paragraphs if p["text"].strip()][:12]
    return result


def search_docx(path: str | Path, query: str, *, match: str = "exact", regex: bool = False) -> list[dict[str, Any]]:
    document = _load(path)
    if not regex:
        return [_span_dict(span) for span in find_text(document, query, match=match)]
    pattern = re.compile(query)
    hits = []
    for index, paragraph in enumerate(document.paragraphs):
        for found in pattern.finditer(paragraph.text):
            hits.append({"text": found.group(0), "paragraph_index": index, "start": found.start(), "context": paragraph.text})
    return hits


def tracked_edit_docx(
    path: str | Path,
    *,
    find: str,
    replace: str,
    nth: int | None = None,
    match: str = "exact",
    author: str | None = None,
) -> dict[str, Any]:
    document = _load(path)
    span = find_one(document, find, nth=nth, match=match)
    result = span.replace(replace, tracked=True, author=author or DEFAULT_AUTHOR)
    save = _save_in_place(path, document)
    return {"replace": result.to_dict(), "revisions": len(_load(path).revisions), "save": save}


def comment_docx(
    path: str | Path,
    *,
    text: str,
    find: str | None = None,
    author: str | None = None,
    nth: int | None = None,
    reply_to: int | None = None,
    resolve_id: int | None = None,
) -> dict[str, Any]:
    document = _load(path)
    author = author or DEFAULT_AUTHOR
    by_id = {comment.comment_id: comment for comment in document.comments}
    done: dict[str, Any] = {}
    if resolve_id is not None:
        if resolve_id not in by_id:
            raise TargetNotFoundError(f"comment {resolve_id} not found; nothing was changed")
        resolve(document, by_id[resolve_id])
        done["resolved"] = resolve_id
    if reply_to is not None:
        if reply_to not in by_id:
            raise TargetNotFoundError(f"comment {reply_to} not found; nothing was changed")
        comment = reply(document, by_id[reply_to], text, author=author)
        done["comment"] = {"comment_id": comment.comment_id, "author": comment.author, "text": comment.text, "parent_id": reply_to}
    elif find:
        span = find_one(document, find, nth=nth)
        comment = span.comment(text, author=author)
        done["comment"] = {"comment_id": comment.comment_id, "author": comment.author, "text": comment.text, "anchored_text": span.text}
    elif resolve_id is None:
        raise ValueError("give find (new comment), reply_to (reply) or resolve (comment id)")
    done["save"] = _save_in_place(path, document)
    return done


def resolve_revisions_docx(
    path: str | Path,
    *,
    action: str,
    scope: str = "all",
    author: str | None = None,
    index: int | None = None,
) -> dict[str, Any]:
    if action not in {"accept", "reject"}:
        raise ValueError("action must be accept or reject")
    document = _load(path)
    revisions = document.revisions
    if scope == "index":
        if index is None:
            raise ValueError("scope=index needs index")
        if not 0 <= index < len(revisions):
            raise TargetNotFoundError(f"revision {index} not found ({len(revisions)} pending); nothing was changed")
        target = revisions[index]
        target.accept() if action == "accept" else target.reject()
        resolved = 1
    else:
        if scope == "author" and not author:
            raise ValueError("scope=author needs author")
        selected = author if scope == "author" else None
        resolved = revisions.accept_all(author=selected) if action == "accept" else revisions.reject_all(author=selected)
    save = _save_in_place(path, document)
    after = _load(path).revisions
    return {"action": action, "scope": scope, "resolved": resolved, "remaining": len(after), "unsupported": after.remaining_unsupported(), "save": save}


def redline_docx(original: str | Path, revised: str | Path, output_path: str | Path, *, author: str | None = None) -> dict[str, Any]:
    result = compare(str(resolve_path(original)), str(resolve_path(revised)), author=author or DEFAULT_AUTHOR)
    out = resolve_path(output_path)
    # The redline is the original package plus tracked changes, so patch against the original.
    patch_save(str(resolve_path(original)), result.document, str(out))
    return {"compare": result.to_dict(), "output_path": str(out)}


def _paragraph_runs(paragraph) -> list[dict[str, Any]]:
    """Read-only walk of a paragraph's runs that keeps tracked insertions and deletions in place."""
    runs: list[dict[str, Any]] = []

    def text_of(run_el, tags: tuple[str, ...]) -> str:
        parts = []
        for node in run_el.iter():
            if node.tag in tags and node.text:
                parts.append(node.text)
            elif node.tag == f"{_W}tab":
                parts.append("\t")
            elif node.tag in (f"{_W}br", f"{_W}cr"):
                parts.append("\n")
        return "".join(parts)

    def push(text: str, rev: str | None, author: str | None) -> None:
        if not text:
            return
        if runs and runs[-1]["rev"] == rev and runs[-1]["author"] == author:
            runs[-1]["text"] += text
        else:
            runs.append({"text": text, "rev": rev, "author": author})

    def walk(parent, rev: str | None = None, author: str | None = None) -> None:
        for child in parent:
            tag = child.tag
            if tag == f"{_W}r":
                push(text_of(child, (f"{_W}t", f"{_W}delText")), rev, author)
            elif tag in (f"{_W}ins", f"{_W}moveTo"):
                walk(child, "ins", child.get(f"{_W}author"))
            elif tag in (f"{_W}del", f"{_W}moveFrom"):
                walk(child, "del", child.get(f"{_W}author"))
            elif tag in (f"{_W}hyperlink", f"{_W}smartTag", f"{_W}fldSimple", f"{_W}sdt", f"{_W}sdtContent"):
                walk(child, rev, author)

    walk(paragraph._p)
    return runs


def preview_docx(path: str | Path, *, max_paragraphs: int = 400) -> dict[str, Any]:
    document = _load(path)
    paragraphs = []
    for paragraph in document.paragraphs[:max_paragraphs]:
        paragraphs.append({"style": paragraph.style.name if paragraph.style is not None else None, "runs": _paragraph_runs(paragraph)})
    comments = [
        {
            "id": entry["comment_id"],
            "author": entry["author"],
            "text": entry["text"],
            "anchor": entry.get("anchored_text"),
            "resolved": entry.get("resolved", False),
            "parent_id": entry.get("parent_id"),
        }
        for entry in comment_thread(document)
    ]
    return {
        "kind": "docx",
        "path": str(resolve_path(path)),
        "docx": {"paragraphs": paragraphs, "comments": comments, "revisions": len(document.revisions), "truncated": len(document.paragraphs) > max_paragraphs},
    }


def _current_text(paragraph) -> str:
    return "".join(run["text"] for run in _paragraph_runs(paragraph) if run["rev"] != "del")


def diff_docx(path: str | Path, baseline: str | Path) -> dict[str, Any]:
    from difflib import SequenceMatcher

    current = _load(path)
    original = _load(baseline)
    before = [_current_text(p) for p in original.paragraphs]
    after = [_current_text(p) for p in current.paragraphs]
    changed = []
    for tag, i1, i2, j1, j2 in SequenceMatcher(a=before, b=after, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        width = max(i2 - i1, j2 - j1)
        for offset in range(width):
            changed.append({
                "before": before[i1 + offset] if i1 + offset < i2 else "",
                "after": after[j1 + offset] if j1 + offset < j2 else "",
            })
    revisions = [
        {"revision_type": r.revision_type, "author": r.author, "text": r.text, "date": r.date.isoformat() if r.date else None}
        for r in current.revisions
    ]
    return {"kind": "docx", "revisions": revisions, "paragraphs_changed": changed}
