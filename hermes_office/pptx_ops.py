from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.diff import diff_decks
from pptx.edit import replace_text, replace_text_at
from pptx.errors import TargetNotFoundError
from pptx.inspect import effective_font, effective_paragraph_format, effective_shape_format, inspect_deck, inspect_text
from pptx.package import patch_save

from .guard import resolve_path

IMPORT_MODES = ("adopt_theme", "keep_appearance", "bake")


def _load(path: str | Path):
    return Presentation(str(resolve_path(path)))


def _save_in_place(path: str | Path, prs) -> dict[str, Any]:
    # patch_save restores original bytes for untouched parts and replaces the file atomically.
    target = str(resolve_path(path))
    residual = patch_save(target, prs, target)
    return residual.to_dict() if hasattr(residual, "to_dict") else {}


def _slide(prs, index: int):
    slides = list(prs.slides)
    if not 0 <= index < len(slides):
        raise TargetNotFoundError(f"slide {index} not found (deck has {len(slides)} slides, 0-based); nothing was changed")
    return slides[index]


def inspect_pptx(path: str | Path, *, detail: str = "summary") -> dict[str, Any]:
    prs = _load(path)
    manifest = inspect_deck(prs).to_dict()
    slides = []
    for index, slide in enumerate(prs.slides):
        blocks = inspect_text(slide).blocks
        texts = [{"shape": block.shape_name, "text": block.text} for block in blocks if block.text.strip()]
        slides.append({
            "index": index,
            "layout": slide.slide_layout.name,
            "shapes": [shape.name for shape in slide.shapes],
            "text": texts if detail == "full" else texts[:6],
        })
    result = {
        "kind": "pptx",
        "path": str(resolve_path(path)),
        "summary": {"slides": manifest["slide_count"], "slide_width": manifest["slide_width"], "slide_height": manifest["slide_height"]},
        "slides": slides,
    }
    if detail == "full":
        result["manifest"] = manifest
    return result


def search_pptx(path: str | Path, query: str, *, match: str = "exact", regex: bool = False) -> list[dict[str, Any]]:
    import re

    prs = _load(path)
    folded = " ".join(query.lower().split())
    pattern = re.compile(query) if regex else None
    hits = []
    for index, slide in enumerate(prs.slides):
        for block in inspect_text(slide).blocks:
            text = block.text
            if pattern is not None:
                found = bool(pattern.search(text))
            elif match == "normalized":
                found = folded in " ".join(text.lower().split())
            else:
                found = query in text
            if found:
                hits.append({"slide": index, "shape": block.shape_name, "text": text, "anchor": block.anchor.to_dict() if hasattr(block.anchor, "to_dict") else None})
    return hits


def inspect_pptx_formats(path: str | Path, *, slide: int, shape: str) -> dict[str, Any]:
    prs = _load(path)
    shp = _slide(prs, slide).shapes.shape_by_name(shape)
    result: dict[str, Any] = {"kind": "pptx", "path": str(resolve_path(path)), "slide": slide, "shape": shape, "shape_format": effective_shape_format(shp).to_dict()}
    if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
        paragraphs = []
        for p_index, paragraph in enumerate(shp.text_frame.paragraphs):
            paragraphs.append({
                "index": p_index,
                "text": "".join(run.text for run in paragraph.runs),
                "paragraph_format": effective_paragraph_format(paragraph).to_dict(),
                "runs": [{"text": run.text, "font": effective_font(run).to_dict()} for run in paragraph.runs],
            })
        result["paragraphs"] = paragraphs
    return result


def edit_pptx_shape(path: str | Path, *, slide: int, ops: list[dict[str, Any]], shape: str | None = None) -> dict[str, Any]:
    prs = _load(path)
    current = _slide(prs, slide)
    applied: list[dict[str, Any]] = []
    for op in ops:
        kind = op.get("op")
        if kind == "replace_text":
            if not shape:
                raise ValueError("replace_text needs shape")
            current.shapes.shape_by_name(shape)  # refuses missing or ambiguous names
            blocks = [b for b in inspect_text(current).blocks if b.shape_name == shape and op["find"] in b.text]
            if not blocks:
                raise TargetNotFoundError(f"{op['find']!r} not found in shape {shape!r} on slide {slide}; nothing was changed")
            count = 0
            for block in blocks:
                count += replace_text_at(prs, block.anchor, op["find"], op["replace"]).replacements
            applied.append({"op": kind, "replacements": count})
        elif kind == "replace_text_deck":
            result = replace_text(prs, op["find"], op["replace"], include_notes=bool(op.get("include_notes", False)))
            if not result.replacements:
                raise TargetNotFoundError(f"{op['find']!r} not found in the deck; nothing was changed")
            applied.append({"op": kind, "replacements": result.replacements})
        elif kind == "set_geometry":
            if not shape:
                raise ValueError("set_geometry needs shape")
            shp = current.shapes.shape_by_name(shape)
            for attr, key in (("left", "x"), ("top", "y"), ("width", "w"), ("height", "h")):
                if op.get(key) is not None:
                    setattr(shp, attr, int(op[key]))
            applied.append({"op": kind, "shape": shape})
        elif kind == "delete_shape":
            if not shape:
                raise ValueError("delete_shape needs shape")
            current.shapes.delete(current.shapes.shape_by_name(shape))
            applied.append({"op": kind, "shape": shape})
        elif kind == "clone_slide":
            clone = prs.slides.clone(current, after=current)
            applied.append({"op": kind, "new_slide_id": clone.slide_id})
        elif kind == "delete_slide":
            prs.slides.delete(current)
            applied.append({"op": kind, "slide": slide})
        elif kind == "move_slide":
            prs.slides.move(current, int(op["to_index"]))
            applied.append({"op": kind, "to_index": int(op["to_index"])})
        else:
            raise ValueError(f"unknown pptx op {kind!r}; use replace_text, replace_text_deck, set_geometry, delete_shape, clone_slide, delete_slide, move_slide")
    save = _save_in_place(path, prs)
    return {"applied": applied, "save": save}


def compose_pptx_decks(base: str | Path, sources: list[dict[str, Any]], *, mode: str, output_path: str | Path) -> dict[str, Any]:
    if mode not in IMPORT_MODES:
        raise ValueError(f"mode must be one of {IMPORT_MODES}")
    base_path = resolve_path(base)
    prs = _load(base_path)
    reports = []
    for source in sources:
        src = _load(source["path"])
        indexes = source.get("slides")
        if indexes is None:
            report = prs.append_deck(src, mode=mode, notes=True)
            reports.append({"source": source["path"], "slides": "all", "report": report.to_dict() if hasattr(report, "to_dict") else None})
            continue
        for index in indexes:
            report = prs.import_slide(src, _slide(src, int(index)), mode=mode, notes=True)
            reports.append({"source": source["path"], "slide": int(index), "report": report.to_dict() if hasattr(report, "to_dict") else None})
    out = resolve_path(output_path)
    patch_save(str(base_path), prs, str(out))
    return {"output_path": str(out), "imports": reports, "slides": len(list(_load(out).slides))}


def diff_pptx(path_a: str | Path, path_b: str | Path, *, detail: str = "text") -> dict[str, Any]:
    return {"kind": "pptx", **diff_decks(str(resolve_path(path_a)), str(resolve_path(path_b)), detail=detail).to_dict()}


def preview_pptx(path: str | Path) -> dict[str, Any]:
    prs = _load(path)
    slides = []
    for index, slide in enumerate(prs.slides):
        shapes = []
        for shape in slide.shapes:
            shapes.append({
                "name": shape.name,
                "kind": shape.shape_type.name if shape.shape_type is not None else "OTHER",
                "text": shape.text_frame.text if shape.has_text_frame else None,
                "x": shape.left,
                "y": shape.top,
                "w": shape.width,
                "h": shape.height,
            })
        title = slide.shapes.title
        slides.append({"index": index, "slide_id": slide.slide_id, "title": title.text_frame.text if title is not None and title.has_text_frame else None, "shapes": shapes})
    return {"kind": "pptx", "path": str(resolve_path(path)), "pptx": {"slides": slides, "slide_w": prs.slide_width, "slide_h": prs.slide_height}}
