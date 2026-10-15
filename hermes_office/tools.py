from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Callable

from docx.errors import PaperRefusal as DocxPaperRefusal
from openpyxl.errors import PaperRefusal as XlsxPaperRefusal
from pptx.errors import PaperRefusal as PptxPaperRefusal

from .bridge import data_dir, emit
from .desk import Desk
from .docx_ops import accept_or_reject_docx, comment_docx, tracked_edit_docx
from .errors import OutsideDeskError, PolicyRefusal, UnsupportedFileError
from .guard import atomic_write, backup, kind_of, refusal_payload, resolve_path, sha256
from .inspect_ops import inspect_file
from .preview import preview_file
from .search_ops import search_path_or_folder
from .diff_ops import diff_file
from .pptx_ops import compose_pptx_decks, diff_pptx, edit_pptx_shape, inspect_pptx_formats
from .xlsx_ops import edit_xlsx, recalc_xlsx, diff_xlsx


def _json(payload: Any) -> str:
    return json.dumps(payload, ensure_ascii=False, default=str)


def _desk() -> Desk:
    return Desk(data_dir())


def _get(args: dict[str, Any], name: str, default: Any = None) -> Any:
    value = args.get(name, default)
    return default if value is None else value


def _workstream_for_write(desk: Desk, args: dict[str, Any], *, path: str | None, tool: str, title: str | None = None) -> dict[str, Any]:
    workstream_id = args.get("workstream_id")
    session_id = args.get("session_id")
    if workstream_id is not None:
        ws = desk.get_workstream(int(workstream_id))
        if ws is None:
            raise OutsideDeskError(f"unknown workstream {workstream_id}")
        return ws
    if session_id:
        ws = desk.workstream_for_session(str(session_id))
        if ws is not None:
            return ws
    if path is None:
        raise OutsideDeskError("a desk path is required")
    folder = desk.folder_for_path(path)
    if folder is None:
        raise OutsideDeskError(str(resolve_path(path)))
    return desk.create_workstream(folder["id"], title or tool, request=args.get("request"), session_id=session_id, status="running")


def _tool_event(desk: Desk, workstream_id: int, kind: str, tool: str, path: str | None, summary: str, op_id: str) -> dict[str, Any]:
    return desk.add_event(workstream_id, kind, tool=tool, path=path, summary=summary, data={"op_id": op_id})


def _refuse(desk: Desk, ws: dict[str, Any] | None, tool: str, path: str | None, exc: Exception, op_id: str | None = None) -> str:
    if ws is not None:
        desk.add_event(ws["id"], "refused", tool=tool, path=path, summary=str(exc), data={"op_id": op_id or tool})
        desk.update_workstream(ws["id"], status="needs_you", needs_you=1, note=str(exc))
        try:
            emit("workstream.updated", {"workstream_id": ws["id"], "status": "needs_you"})
        except Exception:
            pass
    return _json(refusal_payload(exc))


def _finalize(desk: Desk, ws: dict[str, Any], tool: str, path: str | None, op_id: str, receipt: dict[str, Any] | None = None, needs_you: bool = False) -> None:
    desk.add_event(ws["id"], "tool_finished", tool=tool, path=path, summary=tool, data={"op_id": op_id})
    desk.update_workstream(ws["id"], status="needs_you" if needs_you else "done", needs_you=1 if needs_you else 0, note=ws.get("note"))
    if receipt is not None:
        desk.track_file(ws["id"], path or receipt.get("path", ""), "changed", receipt_id=receipt["id"], kind=kind_of(path or receipt.get("path", "")))
    try:
        emit("workstream.updated", {"workstream_id": ws["id"], "status": "needs_you" if needs_you else "done"})
    except Exception:
        pass


def office_desk(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    action = args.get("action")
    try:
        if action == "add_folder":
            path = resolve_path(args["path"])
            if not path.exists() or not path.is_dir():
                raise OutsideDeskError(f"folder does not exist: {path}")
            return _json({"ok": True, "folder": desk.add_folder(path)})
        if action == "remove_folder":
            desk.remove_folder(int(args["folder_id"]))
            return _json({"ok": True})
        if action == "list":
            return _json({"ok": True, "folders": desk.list_folders(), "workstreams": desk.list_workstreams()})
        if action == "status":
            workstream_id = int(args["workstream_id"])
            return _json({"ok": True, "workstream": desk.get_workstream(workstream_id), "now": desk.now(workstream_id)})
        if action == "needs_you":
            ws = desk.update_workstream(int(args["workstream_id"]), status="needs_you", needs_you=1, note=args.get("note"))
            return _json({"ok": True, "workstream": ws})
        if action == "done":
            ws = desk.update_workstream(int(args["workstream_id"]), status="done", needs_you=0, note=args.get("note"))
            return _json({"ok": True, "workstream": ws})
        if action == "attach":
            ws = desk.update_workstream(int(args["workstream_id"]), session_id=args.get("session_id"), title=args.get("title") or desk.get_workstream(int(args["workstream_id"]))["title"])
            return _json({"ok": True, "workstream": ws})
        raise ValueError(f"unknown action: {action}")
    except Exception as exc:
        return _json(refusal_payload(exc))


def office_inspect(args: dict[str, Any], **kwargs: Any) -> str:
    try:
        return _json({"ok": True, **inspect_file(args["path"], detail=args.get("detail", "summary"))})
    except Exception as exc:
        return _json(refusal_payload(exc))


def office_search(args: dict[str, Any], **kwargs: Any) -> str:
    try:
        hits = search_path_or_folder(args.get("path"), args.get("folder"), query=args["query"], match=args.get("match", "exact"), regex=bool(args.get("regex", False)))
        return _json({"ok": True, "matches": hits})
    except Exception as exc:
        return _json(refusal_payload(exc))


def docx_tracked_edit(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    path = args["path"]
    ws = None
    op_id = kwargs.get("task_id") or "docx_tracked_edit"
    try:
        ws = _workstream_for_write(desk, args, path=path, tool="docx_tracked_edit")
        _tool_event(desk, ws["id"], "tool_started", "docx_tracked_edit", path, args["find"], op_id)
        before = sha256(path) if Path(path).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, path) if before is not None else None
        result = tracked_edit_docx(path, find=args["find"], replace=args["replace"], nth=args.get("nth"), match=args.get("match", "exact"), author=args.get("author"))
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "docx_tracked_edit", path, f"replace {args['find']}", before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed", receipt_id=receipt["id"], kind="docx")
        _finalize(desk, ws, "docx_tracked_edit", path, op_id, receipt=receipt, needs_you=True)
        return _json({"ok": True, "receipt": receipt, **result})
    except Exception as exc:
        return _refuse(desk, ws, "docx_tracked_edit", path, exc, op_id)


def docx_comment(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    path = args["path"]
    ws = None
    op_id = kwargs.get("task_id") or "docx_comment"
    try:
        ws = _workstream_for_write(desk, args, path=path, tool="docx_comment")
        _tool_event(desk, ws["id"], "tool_started", "docx_comment", path, args["find"], op_id)
        before = sha256(path) if Path(path).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, path) if before is not None else None
        result = comment_docx(path, find=args["find"], text=args["text"], author=args["author"], nth=args.get("nth"), reply_to=args.get("reply_to"), resolve_comment=bool(args.get("resolve", False)))
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "docx_comment", path, f"comment {args['find']}", before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed", receipt_id=receipt["id"], kind="docx")
        _finalize(desk, ws, "docx_comment", path, op_id, receipt=receipt, needs_you=False)
        return _json({"ok": True, "receipt": receipt, **result})
    except Exception as exc:
        return _refuse(desk, ws, "docx_comment", path, exc, op_id)


def docx_redline_accept_reject(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    path = args["path"]
    ws = None
    op_id = kwargs.get("task_id") or "docx_redline_accept_reject"
    try:
        ws = _workstream_for_write(desk, args, path=path, tool="docx_redline_accept_reject")
        _tool_event(desk, ws["id"], "tool_started", "docx_redline_accept_reject", path, args["action"], op_id)
        before = sha256(path) if Path(path).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, path) if before is not None else None
        result = accept_or_reject_docx(path, args["action"], author=args.get("author"), compare_with=args.get("compare_with"))
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "docx_redline_accept_reject", path, args["action"], before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed", receipt_id=receipt["id"], kind="docx")
        _finalize(desk, ws, "docx_redline_accept_reject", path, op_id, receipt=receipt, needs_you=True)
        return _json({"ok": True, "receipt": receipt, **result})
    except Exception as exc:
        return _refuse(desk, ws, "docx_redline_accept_reject", path, exc, op_id)


def pptx_inspect_formats_handler(args: dict[str, Any], **kwargs: Any) -> str:
    try:
        return _json({"ok": True, **inspect_pptx_formats(args["path"], slide=int(args["slide"]), shape=args["shape"])})
    except Exception as exc:
        return _json(refusal_payload(exc))


def pptx_edit_shape_handler(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    path = args["path"]
    ws = None
    op_id = kwargs.get("task_id") or "pptx_edit_shape"
    try:
        ws = _workstream_for_write(desk, args, path=path, tool="pptx_edit_shape")
        _tool_event(desk, ws["id"], "tool_started", "pptx_edit_shape", path, args["shape"], op_id)
        before = sha256(path) if Path(path).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, path) if before is not None else None
        result = edit_pptx_shape(path, slide=int(args["slide"]), shape=args["shape"], ops=list(args["ops"]))
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "pptx_edit_shape", path, args["shape"], before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed", receipt_id=receipt["id"], kind="pptx")
        _finalize(desk, ws, "pptx_edit_shape", path, op_id, receipt=receipt, needs_you=False)
        return _json({"ok": True, "receipt": receipt, **result})
    except Exception as exc:
        return _refuse(desk, ws, "pptx_edit_shape", path, exc, op_id)


def pptx_compose_decks_handler(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    dest = args["dest"]
    ws = None
    op_id = kwargs.get("task_id") or "pptx_compose_decks"
    try:
        ws = _workstream_for_write(desk, args, path=dest, tool="pptx_compose_decks")
        _tool_event(desk, ws["id"], "tool_started", "pptx_compose_decks", dest, args["mode"], op_id)
        before = sha256(dest) if Path(dest).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, dest) if Path(dest).exists() else None
        result = compose_pptx_decks(dest, list(args["sources"]), mode=args["mode"], output_path=args["output_path"])
        after = sha256(args["output_path"])
        receipt = desk.add_receipt(ws["id"], "pptx_compose_decks", args["output_path"], args["mode"], before, after, backup_path, result)
        desk.track_file(ws["id"], args["output_path"], "created" if before is None else "changed", receipt_id=receipt["id"], kind="pptx")
        _finalize(desk, ws, "pptx_compose_decks", args["output_path"], op_id, receipt=receipt, needs_you=False)
        return _json({"ok": True, "receipt": receipt, **result})
    except Exception as exc:
        return _refuse(desk, ws, "pptx_compose_decks", dest, exc, op_id)


def pptx_diff_handler(args: dict[str, Any], **kwargs: Any) -> str:
    try:
        return _json({"ok": True, "diff": diff_pptx(args["a"], args["b"], detail=args.get("detail", "text"))})
    except Exception as exc:
        return _json(refusal_payload(exc))


def xlsx_edit_handler(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    path = args["path"]
    ws = None
    op_id = kwargs.get("task_id") or "xlsx_edit"
    try:
        ws = _workstream_for_write(desk, args, path=path, tool="xlsx_edit")
        _tool_event(desk, ws["id"], "tool_started", "xlsx_edit", path, "edit", op_id)
        before = sha256(path) if Path(path).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, path) if Path(path).exists() else None
        result = edit_xlsx(path, ops=list(args["ops"]))
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "xlsx_edit", path, "edit", before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed" if before is not None else "created", receipt_id=receipt["id"], kind="xlsx")
        _finalize(desk, ws, "xlsx_edit", path, op_id, receipt=receipt, needs_you=False)
        return _json({"ok": True, "receipt": receipt, **result})
    except Exception as exc:
        return _refuse(desk, ws, "xlsx_edit", path, exc, op_id)


def xlsx_recalc_handler(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    path = args["path"]
    ws = None
    op_id = kwargs.get("task_id") or "xlsx_recalc"
    try:
        ws = _workstream_for_write(desk, args, path=path, tool="xlsx_recalc")
        _tool_event(desk, ws["id"], "tool_started", "xlsx_recalc", path, "recalc", op_id)
        before = sha256(path) if Path(path).exists() else None
        backup_path = backup(desk.data_dir, ws["id"], op_id, path) if Path(path).exists() else None
        result = recalc_xlsx(path, read=list(args["read"]))
        if not result.get("recalculated"):
            receipt = desk.add_receipt(ws["id"], "xlsx_recalc", path, "recalc", before, before, backup_path, result)
            desk.track_file(ws["id"], path, "working", receipt_id=receipt["id"], kind="xlsx")
            _finalize(desk, ws, "xlsx_recalc", path, op_id, receipt=receipt, needs_you=False)
            return _json({"ok": True, **result})
        after = sha256(path)
        receipt = desk.add_receipt(ws["id"], "xlsx_recalc", path, "recalc", before, after, backup_path, result)
        desk.track_file(ws["id"], path, "changed", receipt_id=receipt["id"], kind="xlsx")
        _finalize(desk, ws, "xlsx_recalc", path, op_id, receipt=receipt, needs_you=False)
        return _json({"ok": True, **result})
    except Exception as exc:
        return _refuse(desk, ws, "xlsx_recalc", path, exc, op_id)


def xlsx_diff_handler(args: dict[str, Any], **kwargs: Any) -> str:
    try:
        return _json({"ok": True, "diff": diff_xlsx(args["a"], args["b"])})
    except Exception as exc:
        return _json(refusal_payload(exc))


def office_receipt_handler(args: dict[str, Any], **kwargs: Any) -> str:
    desk = _desk()
    action = args["action"]
    try:
        if action == "list":
            workstream_id = args.get("workstream_id")
            receipts = desk.list_receipts(int(workstream_id)) if workstream_id is not None else desk.list_receipts()
            return _json({"ok": True, "receipts": receipts})
        if action == "show":
            receipt = desk.get_receipt(int(args["receipt_id"]))
            if receipt is None:
                raise OutsideDeskError("receipt not found")
            return _json({"ok": True, "receipt": receipt})
        if action == "revert":
            receipt = desk.get_receipt(int(args["receipt_id"]))
            if receipt is None:
                raise OutsideDeskError("receipt not found")
            path = Path(receipt["path"])
            before = sha256(path) if path.exists() else None
            backup_path = receipt.get("backup_path")
            if backup_path is None:
                raise OutsideDeskError("receipt has no backup")
            revert_backup = backup(desk.data_dir, receipt["workstream_id"], f"revert-{receipt['id']}", path) if path.exists() else None
            src = Path(backup_path)
            atomic_write(path, lambda tmp: tmp.write_bytes(src.read_bytes()))
            after = sha256(path)
            new_receipt = desk.add_receipt(
                receipt["workstream_id"],
                "office_receipt",
                path,
                f"revert receipt {receipt['id']}",
                before,
                after,
                revert_backup or backup_path,
                {"ok": True, "reverted": receipt["id"]},
            )
            desk.mark_reverted(int(args["receipt_id"]))
            desk.track_file(receipt["workstream_id"], path, "changed", receipt_id=new_receipt["id"], kind=kind_of(path))
            return _json({"ok": True, "receipt": new_receipt})
        raise ValueError(f"unknown action: {action}")
    except Exception as exc:
        return _json(refusal_payload(exc))
